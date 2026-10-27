"""
evaluar_agente.py - Evaluation of the LED sales/rental agent: 5 routes -> PDF report.

Each ROUTE is a different customer scenario. For every route the script:
  1. sends the customer input(s) to the REAL agent (same prompt, same MCP tools)
  2. extracts the actual TRAJECTORY (which tools were called, in what order, with what arguments)
  3. checks it against the TRAJECTORY CRITERIA (rule-based, no LLM judge)
  4. checks the RESULT (Excel state + the agent's final reply) against the expected result
  5. writes everything to a PDF (and a JSON file with the raw data)

Safety: the agent runs against a TEMPORARY data folder (LED_DATA_DIR), so your real
leads.xlsx / pricing.xlsx are never touched. Prices come from the default pricing table
created in that temporary folder.

Usage (from the project folder):
  python evaluar_agente.py                    run the 5 routes with the real agent
  python evaluar_agente.py --routes R3,R5     run only some routes
  python evaluar_agente.py --simulate         test the report WITHOUT an LLM (scripted agent)
  python evaluar_agente.py --out informe.pdf  choose the output file

Options: --attempts N (retry a route if the agent crashed, e.g. Gemini 503; default 2),
         --timeout S (seconds per turn; default 240), --pause S (seconds between routes; default 3)
"""
import argparse
import asyncio
import importlib
import json
import os
import platform
import re
import sys
import tempfile
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from trazas import _texto, traza_para_web

HERE = Path(__file__).resolve().parent


# ===========================================================================
# 1. DATA STRUCTURES
# ===========================================================================
@dataclass
class Ctx:
    """Everything a criterion can look at after a route has run."""
    steps: list      # [{"turn": int, "tool": str, "args": dict, "result": any}]
    replies: list    # final agent text of each turn
    estado: dict     # {"lead": row|None, "quotes": [...], "appts": [...]}
    f: dict          # dynamic values (dates)
    oraculo: dict    # expected amounts computed from the pricing table


@dataclass
class Criterio:
    texto: str
    fn: Callable[[Ctx], tuple]   # -> (passed: bool, evidence: str)


@dataclass
class Ruta:
    id: str
    nombre: str
    situacion: str
    turnos: list                 # customer messages, one per turn
    trayectoria_esperada: list   # human-readable expected steps
    resultado_esperado: str
    criterios_trayectoria: list
    criterios_resultado: list


# ===========================================================================
# 2. CRITERIA BUILDERS (small, reusable, rule-based checks)
# ===========================================================================
def _corto(obj, n: int = 150) -> str:
    s = json.dumps(obj, ensure_ascii=False, default=str)
    return s if len(s) <= n else s[: n - 3] + "..."


def _coincide(valor, esperado) -> bool:
    if callable(esperado):
        try:
            return bool(esperado(valor))
        except Exception:
            return False
    if isinstance(esperado, (int, float)):
        try:
            return abs(float(valor) - esperado) < 1e-6
        except (TypeError, ValueError):
            return False
    return str(valor).strip().lower() == str(esperado).strip().lower()


def _llamadas(ctx: Ctx, tool: str, turno=None) -> list:
    return [s for s in ctx.steps if s["tool"] == tool and (turno is None or s["turn"] == turno)]


def _donde(turno) -> str:
    return "" if turno is None else f" in turn {turno + 1}"


def c_primera(texto, tool, **args):
    def fn(ctx):
        if not ctx.steps:
            return False, "No tool was called"
        s = ctx.steps[0]
        ok = s["tool"] == tool and all(_coincide(s["args"].get(k), v) for k, v in args.items())
        return ok, f"First call: {s['tool']}({_corto(s['args'], 110)})"
    return Criterio(texto, fn)


def c_se_llama(texto, tool, turno=None, **args):
    def fn(ctx):
        cand = _llamadas(ctx, tool, turno)
        if not cand:
            return False, f"{tool} was not called{_donde(turno)}"
        for s in cand:
            if all(_coincide(s["args"].get(k), v) for k, v in args.items()):
                return True, f"Found {tool}({_corto(s['args'], 110)})"
        return False, f"{tool} was called but arguments differ: {_corto(cand[0]['args'], 110)}"
    return Criterio(texto, fn)


def c_nunca(texto, tool, turno=None, **args):
    def fn(ctx):
        cand = [s for s in _llamadas(ctx, tool, turno)
                if all(_coincide(s["args"].get(k), v) for k, v in args.items())]
        return (not cand), ("Never called" if not cand else f"Called {len(cand)} time(s){_donde(turno)}")
    return Criterio(texto, fn)


def c_orden(texto, antes, despues):
    def fn(ctx):
        nombres = [s["tool"] for s in ctx.steps]
        if antes not in nombres or despues not in nombres:
            return False, f"Sequence seen: {' -> '.join(nombres) or '(no tools)'}"
        ok = nombres.index(antes) < nombres.index(despues)
        return ok, f"Sequence seen: {' -> '.join(nombres)}"
    return Criterio(texto, fn)


def c_max_llamadas(n):
    def fn(ctx):
        return len(ctx.steps) <= n, f"{len(ctx.steps)} tool call(s)"
    return Criterio(f"At most {n} tool calls (no loops or wasted steps)", fn)


def c_sin_repetidas():
    def fn(ctx):
        vistos = set()
        for s in ctx.steps:
            clave = (s["tool"], json.dumps(s["args"], sort_keys=True, default=str))
            if clave in vistos:
                return False, f"Repeated identical call: {s['tool']}"
            vistos.add(clave)
        return True, "No identical call was repeated"
    return Criterio("The agent does not retry the exact same failing call", fn)


def c_estado(texto, tool, status):
    def fn(ctx):
        for s in _llamadas(ctx, tool):
            r = s["result"] if isinstance(s["result"], dict) else {}
            if r.get("status") == status:
                return True, f"{tool} returned status={status}"
        return False, f"No {tool} call returned status={status}"
    return Criterio(texto, fn)


def c_sin_estado(texto, tool, status):
    def fn(ctx):
        for s in _llamadas(ctx, tool):
            r = s["result"] if isinstance(s["result"], dict) else {}
            if r.get("status") == status:
                return False, f"{tool} returned status={status}"
        return True, f"No {tool} call returned status={status}"
    return Criterio(texto, fn)


# --- result criteria ---------------------------------------------------------
def c_cotizacion(texto, tipo, clave_oraculo):
    def fn(ctx):
        qs = [q for q in ctx.estado["quotes"] if q["type"] == tipo]
        esperado = ctx.oraculo[clave_oraculo]
        if len(qs) != 1:
            return False, f"{len(qs)} {tipo} quotation(s) saved (expected 1)"
        total = float(qs[0]["total"] or 0)
        return abs(total - esperado) < 0.01, f"Saved total {total:,.2f}, expected {esperado:,.2f}"
    return Criterio(texto, fn)


def c_lead_intent(texto, intent):
    def fn(ctx):
        lead = ctx.estado["lead"]
        if not lead:
            return False, "No lead row found in leads.xlsx"
        return lead["intent"] == intent, f"Lead intent saved as '{lead['intent']}'"
    return Criterio(texto, fn)


def c_cita(texto, tipo, prefijo):
    def fn(ctx):
        cs = [a for a in ctx.estado["appts"] if a["type"] == tipo and a["status"] == "scheduled"
              and str(a["datetime"]).startswith(prefijo)]
        return len(cs) == 1, f"{len(cs)} matching appointment(s) saved (expected 1)"
    return Criterio(texto, fn)


def c_sin_citas(texto):
    def fn(ctx):
        n = len(ctx.estado["appts"])
        return n == 0, f"{n} appointment(s) saved"
    return Criterio(texto, fn)


def c_sin_cotizaciones(texto):
    def fn(ctx):
        n = len(ctx.estado["quotes"])
        return n == 0, f"{n} quotation(s) saved"
    return Criterio(texto, fn)


def c_resp_contiene(texto, patron, turno=-1):
    def fn(ctx):
        if not ctx.replies:
            return False, "No reply"
        resp = ctx.replies[turno] if -len(ctx.replies) <= turno < len(ctx.replies) else ""
        ok = bool(re.search(patron, resp, re.I))
        return ok, ("Pattern found in reply" if ok else "Pattern not found in reply")
    return Criterio(texto, fn)


def c_resp_no_contiene(texto, patron):
    def fn(ctx):
        resp = " ".join(ctx.replies)
        ok = not re.search(patron, resp, re.I)
        return ok, ("Not present in reply" if ok else "Found in reply")
    return Criterio(texto, fn)


def c_contacto_soporte(texto):
    def fn(ctx):
        tel = re.sub(r"\D", "", str(ctx.oraculo.get("soporte_phone") or ""))[-10:]
        mail = str(ctx.oraculo.get("soporte_email") or "").lower()
        if not tel and not mail:
            return False, "No support contact is configured in pricing.xlsx > Settings"
        resp = " ".join(ctx.replies)
        ok = bool((mail and mail in resp.lower()) or (tel and tel in re.sub(r"\D", "", resp)))
        return ok, ("Support phone or email found in the reply" if ok
                    else "The reply does not include the support phone or email")
    return Criterio(texto, fn)


def _dims(a) -> list:
    try:
        return sorted([float(a.get("width_m")), float(a.get("height_m"))])
    except (TypeError, ValueError):
        return []


# ===========================================================================
# 3. DYNAMIC VALUES + EXPECTED AMOUNTS (computed independently from the pricing table)
# ===========================================================================
def fechas() -> dict:
    hoy = date.today()
    visita = hoy + timedelta(days=3)
    while visita.weekday() > 4:                 # Monday-Friday only
        visita += timedelta(days=1)
    sabado = hoy + timedelta(days=3)
    while sabado.weekday() != 5:                # next Saturday
        sabado += timedelta(days=1)
    inicio = hoy + timedelta(days=20)
    return {
        "visita_fecha": visita.isoformat(), "visita_dt": f"{visita.isoformat()}T11:00",
        "sabado_fecha": sabado.isoformat(), "sabado_dt": f"{sabado.isoformat()}T12:00",
        "inicio": inicio.isoformat(), "fin": (inicio + timedelta(days=2)).isoformat(),
    }


def _filas(wb, hoja: str) -> list:
    ws = wb[hoja]
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in fila])) for fila in ws.iter_rows(min_row=2)
            if any(c.value is not None for c in fila)]


def _abrir(ruta: Path):
    for intento in range(3):
        try:
            return load_workbook(ruta, data_only=True)
        except Exception:
            if intento == 2:
                raise
            time.sleep(0.3)


def calcular_oraculo(tmp: Path) -> dict:
    """What a correct quotation should cost, computed here from pricing.xlsx
    (independent of the server code, so it can catch pricing bugs)."""
    wb = _abrir(tmp / "pricing.xlsx")
    aj = {r["key"]: r["value"] for r in _filas(wb, "Settings")}
    iva, moneda = float(aj["tax_rate"]), aj["currency"]

    renta = next(r for r in _filas(wb, "Rental") if r["screen_type"] == "outdoor")
    envio = next(r for r in _filas(wb, "Delivery") if str(r["city"]).lower() == "mexico city")
    costo = max(6 * 3 * 3 * float(renta["daily_rate_per_m2"]), float(renta["min_charge"]))
    sub = costo + float(envio["fee"])
    impuesto = round(sub * iva, 2)
    total_renta = round(sub + impuesto, 2)

    prod = next(p for p in _filas(wb, "Purchase")
                if p["screen_type"] == "indoor" and abs(float(p["pixel_pitch_mm"]) - 2.5) < 0.01)
    sub_c = 3 * 2 * 1 * float(prod["price_per_m2"])
    total_compra = round(sub_c + round(sub_c * iva, 2), 2)
    return {"total_renta": total_renta, "total_compra": total_compra,
            "moneda": moneda, "iva": iva,
            "soporte_phone": aj.get("support_phone"), "soporte_email": aj.get("support_email")}


# ===========================================================================
# 4. THE FIVE ROUTES
# ===========================================================================
def construir_rutas(f: dict, o: dict) -> list:
    m = o["moneda"]

    # ---- R1: rental, everything in one message --------------------------------
    def r1_args(ctx):
        for s in _llamadas(ctx, "handle_rental_request"):
            a = s["args"]
            if (str(a.get("screen_type", "")).lower() == "outdoor" and _dims(a) == [3.0, 6.0]
                    and str(a.get("start_date")) == f["inicio"] and str(a.get("end_date")) == f["fin"]):
                return True, f"Found handle_rental_request({_corto(a, 110)})"
        cand = _llamadas(ctx, "handle_rental_request")
        return False, ("handle_rental_request not called" if not cand
                       else f"Arguments differ: {_corto(cand[0]['args'], 110)}")

    r1 = Ruta(
        "R1", "Rental request with complete information",
        "A customer rents an outdoor screen for a 3-day concert and provides every detail "
        "(dates, size, place, contact) in the first message. This is the ideal rental path.",
        [f"Necesito rentar una pantalla LED exterior de 6 x 3 metros para un concierto en Mexico City, "
         f"del {f['inicio']} al {f['fin']}. Soy Ana Ruiz, ana@acme.com, 5512345678, de la empresa Events MX."],
        ["1. evaluate_request(intent=rent) classifies the request and creates the lead.",
         "2. handle_rental_request(screen_type=outdoor, 6 x 3 m, start/end dates, location, contact) "
         "returns status=quotation_created.",
         "3. The agent tells the customer the quotation was prepared (no other tools)."],
        f"A rental quotation is saved for 3 days of an 18 m2 outdoor screen plus the Mexico City delivery "
        f"fee. Expected total: {o['total_renta']:,.2f} {m}. The lead is stored with intent=rent and the "
        f"customer is informed in the reply.",
        [c_primera("First tool call is evaluate_request with intent=rent", "evaluate_request", intent="rent"),
         Criterio("handle_rental_request is called with outdoor, 6 x 3 m and the requested dates", r1_args),
         c_orden("evaluate_request happens before handle_rental_request", "evaluate_request", "handle_rental_request"),
         c_nunca("No purchase tool is used", "handle_purchase_request"),
         c_estado("handle_rental_request ends with status=quotation_created", "handle_rental_request", "quotation_created"),
         c_max_llamadas(3)],
        [c_cotizacion("Exactly one rental quotation is saved with the expected total", "rental", "total_renta"),
         c_lead_intent("The lead is saved with intent=rent", "rent"),
         c_resp_contiene("The reply tells the customer about the quotation", r"cotizaci")],
    )

    # ---- R2: purchase + showroom visit -----------------------------------------
    r2 = Ruta(
        "R2", "Purchase interest and showroom visit booking",
        "A customer wants to buy indoor screens and asks to visit the showroom on a specific weekday at "
        "11:00, giving contact data up front. Tests the appointment path and the availability rules.",
        [f"Quiero comprar pantallas LED interiores y me gustaría visitar el showroom el {f['visita_fecha']} "
         f"a las 11:00. Soy Luis Mora, luis@events.mx, 3398765432."],
        ["1. evaluate_request(intent=buy) classifies the request.",
         f"2. handle_purchase_request(action=showroom_visit, preferred_datetime={f['visita_dt']}) "
         "returns status=scheduled.",
         "3. The agent confirms the date and time to the customer (no quotation, no rental tool)."],
        f"One showroom appointment is stored for {f['visita_fecha']} at 11:00 linked to the lead, the lead "
        "status becomes showroom_visit_booked, and the reply confirms the time.",
        [c_primera("First tool call is evaluate_request with intent=buy", "evaluate_request", intent="buy"),
         c_se_llama("handle_purchase_request books a showroom_visit at the requested date and time",
                    "handle_purchase_request", action="showroom_visit",
                    preferred_datetime=lambda v: str(v).startswith(f["visita_dt"])),
         c_orden("evaluate_request happens before handle_purchase_request", "evaluate_request", "handle_purchase_request"),
         c_nunca("No quotation is requested that the customer did not ask for",
                 "handle_purchase_request", action="send_quotation"),
         c_nunca("No rental tool is used", "handle_rental_request"),
         c_estado("The booking ends with status=scheduled", "handle_purchase_request", "scheduled"),
         c_max_llamadas(3)],
        [c_cita("Exactly one showroom appointment is saved for the requested slot",
                "showroom_visit", f["visita_dt"]),
         c_lead_intent("The lead is saved with intent=buy", "buy"),
         c_resp_contiene("The reply confirms the 11:00 time", r"11[:.]?00|\b11\s?(am|a\.m\.|h)|las 11")],
    )

    # ---- R3: purchase quotation, information arrives over two turns --------------
    def r3_args(ctx):
        for s in _llamadas(ctx, "handle_purchase_request", turno=1):
            a = s["args"]
            if (str(a.get("action")) == "send_quotation"
                    and str(a.get("screen_type", "")).lower() == "indoor"
                    and _coincide(a.get("pixel_pitch_mm"), 2.5) and _dims(a) == [2.0, 3.0]
                    and _coincide(a.get("quantity", 1), 1)):
                return True, f"Found handle_purchase_request({_corto(a, 110)})"
        cand = _llamadas(ctx, "handle_purchase_request", turno=1)
        return False, ("handle_purchase_request not called in turn 2" if not cand
                       else f"Arguments differ: {_corto(cand[0]['args'], 110)}")

    r3 = Ruta(
        "R3", "Vague purchase request completed over two turns",
        "The customer first says only that they want to buy a screen. The agent must NOT act yet. In the "
        "second message the customer gives the specs and contact data and asks for a quotation. Tests "
        "multi-turn memory (the lead_id) and not acting on incomplete information.",
        ["Hola, quiero comprar una pantalla LED.",
         "Quiero una cotización de una pantalla interior pitch 2.5 de 3 x 2 metros, solo una pieza. "
         "Soy Ana Ruiz, ana@acme.com, 5512345678."],
        ["Turn 1: evaluate_request(intent=buy). No action tool yet; the agent offers a call, a showroom "
         "visit or a quotation and asks what the customer prefers.",
         "Turn 2: handle_purchase_request(action=send_quotation, screen_type=indoor, pixel_pitch_mm=2.5, "
         "3 x 2 m, quantity=1) returns status=quotation_created, reusing the lead_id from turn 1.",
         "The agent informs the customer that the quotation was prepared."],
        f"One purchase quotation is saved: 6 m2 of indoor P2.5 plus tax. Expected total: "
        f"{o['total_compra']:,.2f} {m}. The lead is stored with intent=buy.",
        [c_se_llama("Turn 1: evaluate_request classifies the request as buy", "evaluate_request",
                    turno=0, intent="buy"),
         c_nunca("Turn 1: no purchase action is taken with incomplete information",
                 "handle_purchase_request", turno=0),
         Criterio("Turn 2: handle_purchase_request requests a quotation for indoor P2.5, 3 x 2 m, 1 unit", r3_args),
         c_estado("The quotation ends with status=quotation_created", "handle_purchase_request", "quotation_created"),
         c_nunca("No rental tool is used", "handle_rental_request"),
         c_max_llamadas(4)],
        [c_cotizacion("Exactly one purchase quotation is saved with the expected total", "purchase", "total_compra"),
         c_lead_intent("The lead is saved with intent=buy", "buy"),
         c_resp_contiene("Turn 1 reply offers the options (call, showroom visit or quotation)",
                         r"llamada|visita|showroom|cotizaci", turno=0),
         c_resp_contiene("Final reply tells the customer about the quotation", r"cotizaci")],
    )

    # ---- R4: unsupported request -------------------------------------------------
    r4 = Ruta(
        "R4", "Unsupported request (repair instead of buy or rent)",
        "An existing customer asks for a repair and its cost. The company only handles buying and renting, "
        "so the agent must decline politely without inventing prices or promises.",
        ["Hola, se dañó una pantalla LED que compré con ustedes hace un año. ¿Pueden repararla y cuánto costaría?"],
        ["1. evaluate_request(intent=other) classifies the request as unsupported, logs the lead and "
         "returns the support-team contact.",
         "2. No purchase or rental tool is called.",
         "3. The agent explains this chat only handles buying and renting and gives the customer the "
         "support-team phone or email taken from the tool result (it invents no contact data)."],
        f"The lead is logged as unsupported. No quotation and no appointment are created. The reply declines "
        f"politely, invents no repair price or time, and gives the support contact "
        f"({o.get('soporte_phone')} / {o.get('soporte_email')}) so the customer can call or write.",
        [c_se_llama("evaluate_request classifies the request as unsupported (intent=other)",
                    "evaluate_request", intent="other"),
         c_nunca("No purchase tool is used", "handle_purchase_request"),
         c_nunca("No rental tool is used", "handle_rental_request"),
         c_max_llamadas(2)],
        [c_lead_intent("The lead is logged with intent=unsupported", "unsupported"),
         c_sin_cotizaciones("No quotation is created"),
         c_sin_citas("No appointment is created"),
         c_resp_no_contiene("The reply does not invent a price",
                            r"\$\s*\d|\d[\d,.]*\s*(MXN|pesos|USD|d[oó]lares)"),
         c_contacto_soporte("The reply gives the support team's phone or email from the tool result")],
    )

    # ---- R5: error recovery (closed day) --------------------------------------------
    r5 = Ruta(
        "R5", "Booking rejected by a business rule (showroom closed on Saturday)",
        "The customer asks for a showroom visit on a Saturday, when the showroom is closed. The tool rejects "
        "the booking. The agent must recover: not insist, not book a different day on its own, and ask the "
        "customer for another day.",
        [f"Quiero comprar pantallas y visitar el showroom el sábado {f['sabado_fecha']} a las 12:00. "
         "Soy Pedro Gil, pedro@mail.com, 5500000000."],
        ["1. evaluate_request(intent=buy).",
         f"2. handle_purchase_request(action=showroom_visit, preferred_datetime={f['sabado_dt']}) returns an "
         "error saying the showroom is closed that day.",
         "3. The agent stops trying, books nothing, and asks the customer for another day or time."],
        "No appointment exists for this lead. The reply explains the day is not available and asks the "
        "customer to choose another date.",
        [c_primera("First tool call is evaluate_request with intent=buy", "evaluate_request", intent="buy"),
         c_se_llama("The agent attempts the requested Saturday slot",
                    "handle_purchase_request", action="showroom_visit",
                    preferred_datetime=lambda v: str(v).startswith(f["sabado_fecha"])),
         c_sin_estado("No booking is confirmed on a closed day", "handle_purchase_request", "scheduled"),
         c_sin_repetidas(),
         c_max_llamadas(4)],
        [c_sin_citas("No appointment is saved"),
         c_resp_contiene("The reply asks the customer for another day or time",
                         r"otro|otra|diferente|alternativ|distint|elegir|prefier"),
         c_resp_contiene("The reply explains the day is not available",
                         r"cerrad|no (abre|atiende|est[aá] disponible|hay)|lunes a viernes|no disponible|fines? de semana")],
    )
    return [r1, r2, r3, r4, r5]


# ===========================================================================
# 5. AGENTS: real (LangChain + MCP) or scripted (for --simulate)
# ===========================================================================
async def crear_agente_real(tmp: Path):
    from langchain.agents import create_agent
    from langchain_mcp_adapters.client import MultiServerMCPClient
    from config import PYTHON, llm_langchain

    mod = importlib.import_module("07_mi_agente")        # same prompt + server list as the chatbot
    entorno = {**os.environ, "LED_DATA_DIR": str(tmp)}   # the MCP subprocess must use the temp folder
    cliente = MultiServerMCPClient({
        nombre: {"command": PYTHON, "args": [ruta], "transport": "stdio", "env": entorno}
        for nombre, ruta in mod.SERVIDORES.items()
    })
    return create_agent(model=llm_langchain(), tools=await cliente.get_tools(),
                        system_prompt=mod.SYSTEM_PROMPT)


class _Msg:
    def __init__(self, tipo, contenido, tool_calls=None, tool_call_id=None):
        self.type, self.content = tipo, contenido
        self.tool_calls = tool_calls or []
        self.tool_call_id = tool_call_id


class _Seguro(dict):
    def __missing__(self, k):
        return "{" + k + "}"


class AgenteSimulado:
    """A scripted 'good agent' that calls the REAL tool functions. It only exists to test
    the report/criteria without an LLM. Its output is NOT evidence about the real agent."""

    def __init__(self, srv, f):
        self.srv, self.f, self.ruta, self.lead = srv, f, None, None
        c = {"name": "Ana Ruiz", "email": "ana@acme.com", "phone": "5512345678"}
        self.planes = {
            "R1": [{"llamadas": [
                ("evaluate_request", {"intent": "rent", "confidence": 0.95, "summary": "Renta para concierto"}),
                ("handle_rental_request", {"lead_id": "$lead", "start_date": f["inicio"], "end_date": f["fin"],
                                           "location": "Mexico City", "screen_type": "outdoor",
                                           "width_m": 6, "height_m": 3, "company": "Events MX", **c})],
                "respuesta": "Listo, Ana. Te envié la cotización de renta. Total: {total} MXN."}],
            "R2": [{"llamadas": [
                ("evaluate_request", {"intent": "buy", "confidence": 0.9, "summary": "Compra y visita"}),
                ("handle_purchase_request", {"lead_id": "$lead", "action": "showroom_visit",
                                             "preferred_datetime": f["visita_dt"], "name": "Luis Mora",
                                             "email": "luis@events.mx", "phone": "3398765432"})],
                "respuesta": "Tu visita al showroom quedó agendada para el {visita_fecha} a las 11:00."}],
            "R3": [{"llamadas": [("evaluate_request", {"intent": "buy", "confidence": 0.9, "summary": "Compra"})],
                    "respuesta": "¡Con gusto! ¿Prefieres una llamada, una visita al showroom o una cotización?"},
                   {"llamadas": [("handle_purchase_request", {"lead_id": "$lead", "action": "send_quotation",
                                                              "screen_type": "indoor", "pixel_pitch_mm": 2.5,
                                                              "width_m": 3, "height_m": 2, "quantity": 1, **c})],
                    "respuesta": "Te envié la cotización de compra. Total: {total} MXN."}],
            "R4": [{"llamadas": [("evaluate_request", {"intent": "other", "confidence": 0.95,
                                                       "summary": "Reparación de pantalla"})],
                    "respuesta": "Por ahora solo puedo ayudarte con la compra o renta de pantallas LED. "
                                 "Para reparaciones, llama a nuestro equipo de soporte al {telefono} o escribe a {correo}."}],
            "R5": [{"llamadas": [
                ("evaluate_request", {"intent": "buy", "confidence": 0.9, "summary": "Visita en sábado"}),
                ("handle_purchase_request", {"lead_id": "$lead", "action": "showroom_visit",
                                             "preferred_datetime": f["sabado_dt"], "name": "Pedro Gil",
                                             "email": "pedro@mail.com", "phone": "5500000000"})],
                "respuesta": "El showroom está cerrado ese día. ¿Podrías elegir otro día entre lunes y viernes?"}],
        }

    async def ainvoke(self, entrada):
        msgs = list(entrada["messages"])
        es_humano = lambda m: (getattr(m, "type", None) == "human") or (isinstance(m, dict) and m.get("role") == "user")
        turno = sum(1 for m in msgs if es_humano(m)) - 1
        if turno == 0:
            self.lead = None
        plan = self.planes[self.ruta][turno]
        previos = [m if hasattr(m, "type") else _Msg("human", m["content"]) for m in msgs]
        nuevos, ultimo = [], {}
        for i, (tool, args) in enumerate(plan["llamadas"]):
            args = {k: (self.lead if v == "$lead" else v) for k, v in args.items()}
            res = await asyncio.to_thread(getattr(self.srv, tool), **args)
            if isinstance(res, dict):
                ultimo = res
                if res.get("lead_id"):
                    self.lead = res["lead_id"]
            cid = f"call_{turno}_{i}"
            nuevos += [_Msg("ai", "", [{"name": tool, "args": args, "id": cid}]),
                       _Msg("tool", json.dumps(res, default=str), tool_call_id=cid)]
        ctx = _Seguro({**self.f, **{k: v for k, v in ultimo.items()}})
        if "total" in ultimo:
            ctx["total"] = f"{float(ultimo['total']):,.2f}"
        soporte = ultimo.get("support_contact") or {}
        ctx["telefono"], ctx["correo"] = soporte.get("phone", ""), soporte.get("email", "")
        nuevos.append(_Msg("ai", plan["respuesta"].format_map(ctx)))
        return {"messages": previos + nuevos}


# ===========================================================================
# 6. RUNNING A ROUTE
# ===========================================================================
def leer_estado(tmp: Path, lead_id) -> dict:
    vacio = {"lead": None, "quotes": [], "appts": []}
    if not lead_id or not (tmp / "leads.xlsx").exists():
        return vacio
    wb = _abrir(tmp / "leads.xlsx")
    return {
        "lead": next((r for r in _filas(wb, "Leads") if r["lead_id"] == lead_id), None),
        "quotes": [r for r in _filas(wb, "Quotations") if r["lead_id"] == lead_id],
        "appts": [r for r in _filas(wb, "Appointments") if r["lead_id"] == lead_id],
    }


def _lead_de(steps):
    lead = None
    for s in steps:
        if isinstance(s["result"], dict) and s["result"].get("lead_id"):
            lead = s["result"]["lead_id"]
    return lead


def _evaluar(criterios: list, ctx: Ctx) -> list:
    salida = []
    for c in criterios:
        try:
            ok, ev = c.fn(ctx)
        except Exception as e:
            ok, ev = False, f"Check error: {type(e).__name__}: {e}"
        salida.append({"criterio": c.texto, "ok": bool(ok), "evidencia": ev})
    return salida


async def ejecutar_turnos(ruta: Ruta, agente, timeout: float):
    historial, steps, replies, errores, duraciones = [], [], [], [], []
    if hasattr(agente, "ruta"):
        agente.ruta = ruta.id
    for i, texto in enumerate(ruta.turnos):
        t0 = time.time()
        try:
            r = await asyncio.wait_for(
                agente.ainvoke({"messages": historial + [{"role": "user", "content": texto}]}), timeout)
            mensajes = r["messages"]
            for p in traza_para_web(mensajes[len(historial) + 1:]):
                steps.append({**p, "turn": i})
            replies.append(_texto(mensajes[-1].content).strip())
            historial = mensajes
        except Exception as e:
            errores.append(f"Turn {i + 1}: {type(e).__name__}: {str(e)[:300]}")
            replies.append("")
            break
        finally:
            duraciones.append(round(time.time() - t0, 1))
    return steps, replies, errores, duraciones


async def correr_ruta(ruta: Ruta, agente, tmp: Path, f: dict, o: dict, intentos: int, timeout: float) -> dict:
    for intento in range(1, intentos + 1):
        steps, replies, errores, duraciones = await ejecutar_turnos(ruta, agente, timeout)
        if not errores:
            break
        if intento < intentos:
            print(f"   attempt {intento} failed ({errores[-1][:90]}); retrying in 20 s ...")
            await asyncio.sleep(20)
    lead_id = _lead_de(steps)
    ctx = Ctx(steps, replies, leer_estado(tmp, lead_id), f, o)
    tray, res = _evaluar(ruta.criterios_trayectoria, ctx), _evaluar(ruta.criterios_resultado, ctx)
    aviso = None
    if lead_id and ctx.estado["lead"] is None:
        aviso = ("The lead was not found in the temporary data folder. The MCP server may be using a fixed "
                 "data path, so isolation is NOT confirmed.")
    return {
        "id": ruta.id, "nombre": ruta.nombre, "situacion": ruta.situacion,
        "turnos": ruta.turnos, "trayectoria_esperada": ruta.trayectoria_esperada,
        "resultado_esperado": ruta.resultado_esperado,
        "criterios_trayectoria": tray, "criterios_resultado": res,
        "trayectoria_real": steps, "respuestas": replies, "errores": errores,
        "duraciones_s": duraciones, "intentos": intento, "lead_id": lead_id, "aviso": aviso,
        "ok_trayectoria": all(c["ok"] for c in tray) and bool(tray),
        "ok_resultado": all(c["ok"] for c in res) and bool(res),
    }


# ===========================================================================
# 7. PDF REPORT
# ===========================================================================
def _resumen_resultado(r) -> str:
    if isinstance(r, dict):
        if "error" in r:
            return "ERROR: " + str(r["error"])[:170]
        partes = []
        for k in ("status", "lead_id", "quote_id", "appointment_id", "when"):
            if k in r:
                partes.append(f"{k}={r[k]}")
        if "total" in r:
            partes.append(f"total={r['total']} {r.get('currency', '')}")
        return ", ".join(partes) or _corto(r, 170)
    return str(r)[:170] if r is not None else "(no result)"


def generar_pdf(resultados: list, meta: dict, destino: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    NAVY, GREEN, RED = colors.HexColor("#22407A"), colors.HexColor("#1E7A46"), colors.HexColor("#B3261E")
    GREY, LINE = colors.HexColor("#F1F3F8"), colors.HexColor("#D3D9E4")
    ancho = A4[0] - 36 * mm

    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9, leading=11.5, textColor=colors.HexColor("#1B2333"))
    peq = ParagraphStyle("peq", parent=base, fontSize=8, leading=10.5)
    mono = ParagraphStyle("mono", parent=base, fontName="Courier", fontSize=7.8, leading=10, wordWrap="CJK")
    h_titulo = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=20, leading=24)
    h_sec = ParagraphStyle("s", parent=base, fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                           textColor=NAVY, spaceBefore=6, spaceAfter=2)
    blanco = ParagraphStyle("w", parent=base, fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.white)

    def P(t, st=base):
        return Paragraph(escape(str(t)), st)

    def veredicto(ok, texto_ok="PASS", texto_no="FAIL"):
        col = GREEN if ok else RED
        return Paragraph(f'<font color="{col.hexval().replace("0x", "#")}"><b>{texto_ok if ok else texto_no}</b></font>', base)

    def tabla(filas, anchos, cab=True):
        t = Table(filas, colWidths=anchos, repeatRows=1 if cab else 0)
        est = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), 0.4, LINE),
               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if cab:
            est += [("BACKGROUND", (0, 0), (-1, 0), GREY)]
        t.setStyle(TableStyle(est))
        return t

    def pie(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#5A6478"))
        etiqueta = "SIMULATED RUN - not real agent output" if meta["simulado"] else "LED agent evaluation"
        canvas.drawString(18 * mm, 10 * mm, f"{etiqueta}  |  {meta['fecha']}")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    def bloque_criterios(titulo, lista):
        filas = [[P("#", peq), P("Criterion", peq), P("Verdict", peq), P("Evidence", peq)]]
        for i, c in enumerate(lista, 1):
            filas.append([P(i, peq), P(c["criterio"], peq), veredicto(c["ok"]), P(c["evidencia"], peq)])
        return KeepTogether([P(titulo, h_sec), tabla(filas, [8 * mm, 66 * mm, 16 * mm, ancho - 90 * mm])])

    historia = []

    # ---------------- cover + summary ----------------
    historia.append(P("Agent evaluation report", h_titulo))
    historia.append(P("LED screens sales and rental assistant - 5 evaluation routes", base))
    historia.append(Spacer(1, 6))
    if meta["simulado"]:
        banner = Table([[Paragraph("<b>SIMULATED RUN.</b> This report was produced with a scripted agent only to "
                                   "test the report format. It says nothing about the real agent.", base)]],
                       colWidths=[ancho])
        banner.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FDECEA")),
                                    ("BOX", (0, 0), (-1, -1), 0.8, RED), ("PADDING", (0, 0), (-1, -1), 6)]))
        historia += [banner, Spacer(1, 8)]

    info = [[P("Date", peq), P(meta["fecha"], peq)],
            [P("Agent", peq), P(meta["agente"], peq)],
            [P("Environment", peq), P(meta["entorno"], peq)],
            [P("Data", peq), P("Isolated temporary data folder; your real Excel files were not used.", peq)]]
    historia += [tabla(info, [28 * mm, ancho - 28 * mm], cab=False), Spacer(1, 8)]

    total_tray = sum(len(r["criterios_trayectoria"]) for r in resultados)
    ok_tray = sum(c["ok"] for r in resultados for c in r["criterios_trayectoria"])
    total_res = sum(len(r["criterios_resultado"]) for r in resultados)
    ok_res = sum(c["ok"] for r in resultados for c in r["criterios_resultado"])
    pasadas = sum(r["ok_trayectoria"] and r["ok_resultado"] for r in resultados)

    historia.append(P("Summary", h_sec))
    historia.append(Paragraph(
        f"<b>{pasadas} of {len(resultados)} routes passed.</b> Trajectory criteria met: {ok_tray}/{total_tray}. "
        f"Result criteria met: {ok_res}/{total_res}.", base))
    historia.append(Spacer(1, 6))
    filas = [[P("Route", peq), P("Scenario", peq), P("Trajectory", peq), P("Result", peq), P("Verdict", peq)]]
    for r in resultados:
        nt, nr = len(r["criterios_trayectoria"]), len(r["criterios_resultado"])
        filas.append([P(r["id"], peq), P(r["nombre"], peq),
                      P(f"{sum(c['ok'] for c in r['criterios_trayectoria'])}/{nt}", peq),
                      P(f"{sum(c['ok'] for c in r['criterios_resultado'])}/{nr}", peq),
                      veredicto(r["ok_trayectoria"] and r["ok_resultado"])])
    historia.append(tabla(filas, [14 * mm, ancho - 14 * mm - 22 * mm - 22 * mm - 20 * mm, 22 * mm, 22 * mm, 20 * mm]))

    historia.append(P("How to read this report", h_sec))
    for t in ["Each route is a different customer scenario sent to the real agent.",
              "Expected trajectory: the tool calls the agent should make, in order.",
              "Trajectory criteria: rule-based checks over the real tool calls (order, arguments, forbidden "
              "tools, no loops). No LLM is used as a judge.",
              "Result trajectory: the tool calls the agent actually made, with their outcomes.",
              "Expected result: what must be true afterwards (Excel records and the customer reply), checked "
              "with the result criteria listed under it."]:
        historia.append(Paragraph("- " + escape(t), base))
    if meta["avisos"]:
        historia.append(P("Warnings", h_sec))
        for a in meta["avisos"]:
            historia.append(P("- " + a, base))

    # ---------------- one section per route ----------------
    for r in resultados:
        historia.append(PageBreak())
        ok = r["ok_trayectoria"] and r["ok_resultado"]
        col = GREEN if ok else RED
        barra = Table([[Paragraph(escape(f"Route {r['id'][1:]} - {r['nombre']}"), blanco),
                        Paragraph(f'<b>{"PASS" if ok else "FAIL"}</b>', blanco)]],
                      colWidths=[ancho - 24 * mm, 24 * mm])
        barra.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, 0), NAVY), ("BACKGROUND", (1, 0), (1, 0), col),
                                   ("ALIGN", (1, 0), (1, 0), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                   ("PADDING", (0, 0), (-1, -1), 7)]))
        historia += [barra, Spacer(1, 4), P(r["situacion"], base)]

        historia.append(P("Input", h_sec))
        for i, t in enumerate(r["turnos"]):
            caja = Table([[P(f"Turn {i + 1} - customer:", peq)], [P(t, base)]], colWidths=[ancho])
            caja.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), GREY), ("BOX", (0, 0), (-1, -1), 0.4, LINE),
                                      ("PADDING", (0, 0), (-1, -1), 5)]))
            historia += [caja, Spacer(1, 3)]

        historia.append(P("Expected trajectory", h_sec))
        for linea in r["trayectoria_esperada"]:
            historia.append(P(linea, base))

        historia.append(P("Expected result", h_sec))
        historia.append(P(r["resultado_esperado"], base))
        historia.append(bloque_criterios("Result criteria (checked after the run)", r["criterios_resultado"]))

        historia.append(bloque_criterios("Trajectory criteria", r["criterios_trayectoria"]))

        historia.append(P("Result trajectory (what the agent actually did)", h_sec))
        if r["trayectoria_real"]:
            filas = [[P("#", peq), P("Turn", peq), P("Tool", peq), P("Arguments", peq), P("Outcome", peq)]]
            for i, s in enumerate(r["trayectoria_real"], 1):
                filas.append([P(i, peq), P(s["turn"] + 1, peq), P(s["tool"], peq),
                              Paragraph(escape(_corto(s["args"], 200)), mono),
                              P(_resumen_resultado(s["result"]), peq)])
            historia.append(tabla(filas, [7 * mm, 11 * mm, 41 * mm, 63 * mm, ancho - 122 * mm]))
        else:
            historia.append(P("The agent called no tools.", base))
        historia.append(Spacer(1, 4))
        for i, resp in enumerate(r["respuestas"]):
            caja = Table([[P(f"Agent reply, turn {i + 1}:", peq)], [P(resp or "(no reply)", base)]], colWidths=[ancho])
            caja.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.white), ("BOX", (0, 0), (-1, -1), 0.4, LINE),
                                      ("PADDING", (0, 0), (-1, -1), 5)]))
            historia += [caja, Spacer(1, 3)]
        if r["errores"]:
            historia.append(P("Run errors: " + " | ".join(r["errores"]), peq))
        if r["aviso"]:
            historia.append(P("Warning: " + r["aviso"], peq))

        historia.append(Spacer(1, 4))
        historia.append(P(f"Attempts: {r['intentos']}   |   Time per turn (s): "
                          f"{', '.join(str(d) for d in r['duraciones_s'])}   |   Lead: {r['lead_id'] or 'none'}", peq))

    doc = SimpleDocTemplate(str(destino), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=13 * mm, bottomMargin=16 * mm, title="Agent evaluation report",
                            author="LED sales assistant project")
    doc.build(historia, onFirstPage=pie, onLaterPages=pie)


# ===========================================================================
# 8. MAIN
# ===========================================================================
def verificar_aislamiento() -> None:
    src = HERE / "servidor_led.py"
    if src.exists() and "LED_DATA_DIR" not in src.read_text(encoding="utf-8"):
        sys.exit("servidor_led.py does not read the LED_DATA_DIR environment variable, so the evaluation could "
                 "write test rows into your REAL Excel files. Keep this line in servidor_led.py:\n"
                 '  DATA_DIR = Path(os.getenv("LED_DATA_DIR", "<your folder>"))\nand run again.')


async def principal(args) -> None:
    ahora = datetime.now()
    tmp = Path(tempfile.mkdtemp(prefix="eval_led_"))
    os.environ["LED_DATA_DIR"] = str(tmp)           # set BEFORE importing servidor_led
    sys.path.insert(0, str(HERE))
    if not args.simulate:
        verificar_aislamiento()

    import servidor_led as srv
    srv.ensure_files()
    f, o = fechas(), None
    o = calcular_oraculo(tmp)
    rutas = construir_rutas(f, o)
    if args.routes:
        pedidas = {x.strip().upper() for x in args.routes.split(",")}
        rutas = [r for r in rutas if r.id in pedidas]
        if not rutas:
            sys.exit("No matching routes. Use ids like R1,R3,R5")

    if args.simulate:
        agente, desc_agente = AgenteSimulado(srv, f), "Scripted agent (simulation, no LLM)"
    else:
        print("Building the real agent (LLM + MCP tools) ...")
        agente = await crear_agente_real(tmp)
        import config
        modelo = config.GOOGLE_MODEL if config.PROVEEDOR == "google" else config.OLLAMA_MODEL
        desc_agente = f"LangChain agent, provider={config.PROVEEDOR}, model={modelo}"

    resultados = []
    for i, ruta in enumerate(rutas):
        print(f"[{ruta.id}] {ruta.nombre} ...")
        r = await correr_ruta(ruta, agente, tmp, f, o, args.attempts, args.timeout)
        resultados.append(r)
        print(f"     trajectory {'PASS' if r['ok_trayectoria'] else 'FAIL'} | "
              f"result {'PASS' if r['ok_resultado'] else 'FAIL'}"
              + (f" | errors: {r['errores']}" if r["errores"] else ""))
        if i < len(rutas) - 1 and not args.simulate:
            await asyncio.sleep(args.pause)

    avisos = sorted({r["aviso"] for r in resultados if r["aviso"]})
    meta = {"fecha": ahora.strftime("%Y-%m-%d %H:%M"), "agente": desc_agente, "simulado": args.simulate,
            "entorno": f"Python {platform.python_version()} on {platform.system()}", "avisos": avisos}

    base = f"evaluacion_agente_{'SIMULADA_' if args.simulate else ''}{ahora:%Y%m%d_%H%M}"
    destino = Path(args.out) if args.out else HERE / f"{base}.pdf"
    generar_pdf(resultados, meta, destino)
    destino.with_suffix(".json").write_text(
        json.dumps({"meta": meta, "rutas": resultados}, indent=2, ensure_ascii=False, default=str), encoding="utf-8")

    pasadas = sum(r["ok_trayectoria"] and r["ok_resultado"] for r in resultados)
    print(f"\n{pasadas}/{len(resultados)} routes passed.\nPDF:  {destino}\nJSON: {destino.with_suffix('.json')}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Evaluate the LED agent on 5 routes and write a PDF report.")
    ap.add_argument("--simulate", action="store_true", help="use a scripted agent (tests the report, no LLM)")
    ap.add_argument("--routes", default="", help="comma-separated route ids, e.g. R1,R3")
    ap.add_argument("--out", default="", help="output PDF path")
    ap.add_argument("--attempts", type=int, default=2, help="attempts per route if the agent crashes")
    ap.add_argument("--timeout", type=float, default=240, help="seconds allowed per turn")
    ap.add_argument("--pause", type=float, default=3, help="seconds to wait between routes")
    asyncio.run(principal(ap.parse_args()))
