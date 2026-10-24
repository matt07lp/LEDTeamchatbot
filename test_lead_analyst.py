"""
Tests for lead_analyst.py (agent 2). They use a FAKE LLM and temporary Excel files,
so no API key is needed.

Run:  pytest test_lead_analyst.py -v
"""
import asyncio

import pytest
from openpyxl import load_workbook

import lead_analyst as la

LEAD = "L-TEST0001"


class Msg:
    def __init__(self, type, content):
        self.type = type
        self.content = content


class FakeLLM:
    """Mimics llm.with_structured_output(Schema).ainvoke(messages)."""
    def __init__(self, senales=None, falla=False):
        self.senales, self.falla, self.recibido = senales, falla, None

    def with_structured_output(self, schema):
        return self

    async def ainvoke(self, mensajes):
        self.recibido = mensajes
        if self.falla:
            raise RuntimeError("llm down")
        return self.senales


SENALES_FUERTES = la.LeadSignals(
    urgency="high", decision_maker="decision_maker", requirements_clarity="high",
    budget_signal="high", reasons=["Evento en 2 semanas", "Es el dueño"],
    next_action="Llamar hoy")


@pytest.fixture
def datos(tmp_path, monkeypatch):
    """Temp data folder with one lead that booked a showroom visit and got a quotation."""
    pytest.importorskip("mcp")
    import servidor_led as srv
    monkeypatch.setattr(srv, "DATA_DIR", tmp_path)
    monkeypatch.setattr(srv, "LEADS_FILE", tmp_path / "leads.xlsx")
    monkeypatch.setattr(srv, "PRICING_FILE", tmp_path / "pricing.xlsx")
    srv.ensure_files()
    wb = load_workbook(srv.LEADS_FILE)
    wb["Leads"].append([LEAD, "2026-10-08T10:00:00", "2026-10-08T10:00:00", "buy",
                        "Wants 4 outdoor screens", "Ana Ruiz", "ana@acme.com",
                        "5512345678", "Acme", "purchase_quoted"])
    wb["Appointments"].append(["A-1", LEAD, "showroom_visit", "2026-10-20T15:00:00", "scheduled"])
    wb["Quotations"].append(["Q-1", LEAD, "purchase", "{}", 100000, 16000, 116000,
                             "MXN", "2026-10-25", "emailed"])
    wb.save(srv.LEADS_FILE)
    return tmp_path


def run(coro):
    return asyncio.run(coro)


# --- pure scoring -------------------------------------------------------------
def hechos(accion="none", total=0.0, empresa=""):
    return {"best_action": accion, "deal_total": total, "company": empresa}


def senales(**kw):
    base = dict(urgency="unknown", decision_maker="unknown",
                requirements_clarity="low", budget_signal="unknown")
    return {**base, **kw}


def test_strong_lead_is_hot():
    r = la.reglas_por_defecto()
    s = senales(urgency="high", decision_maker="decision_maker",
                requirements_clarity="high", budget_signal="high")
    score, desglose = la.calcular_puntaje(s, hechos("showroom_visit", 120000, "Acme"), r)
    assert score == 95 and la.prioridad_de(score, r) == "HOT"
    assert sum(desglose.values()) == 95


def test_weak_lead_is_cold():
    r = la.reglas_por_defecto()
    score, _ = la.calcular_puntaje(senales(urgency="low", decision_maker="researching"),
                                   hechos("quotation", 5000), r)
    assert score == 10 and la.prioridad_de(score, r) == "COLD"


def test_priority_thresholds_boundaries():
    r = la.reglas_por_defecto()
    assert la.prioridad_de(65, r) == "HOT"
    assert la.prioridad_de(64, r) == "WARM"
    assert la.prioridad_de(35, r) == "WARM"
    assert la.prioridad_de(34, r) == "COLD"


def test_score_is_clamped_to_100():
    r = la.reglas_por_defecto()
    r["weights"][("urgency", "high")] = 500
    score, _ = la.calcular_puntaje(senales(urgency="high"), hechos(), r)
    assert score == 100


def test_deal_size_uses_highest_matching_tier():
    r = la.reglas_por_defecto()
    _, d = la.calcular_puntaje(senales(), hechos(total=150000), r)
    assert d["deal_size"] == 10
    _, d = la.calcular_puntaje(senales(), hechos(total=300000), r)
    assert d["deal_size"] == 15


# --- facts from Excel ---------------------------------------------------------
def test_facts_pick_best_action_and_total(datos):
    h = la.reunir_hechos(LEAD)
    assert h["best_action"] == "showroom_visit"
    assert h["deal_total"] == 116000 and h["currency"] == "MXN"
    assert h["company"] == "Acme"


def test_facts_unknown_lead_raises(datos):
    with pytest.raises(ValueError):
        la.reunir_hechos("L-NOPE0000")


# --- coordination helpers -------------------------------------------------------
def test_lead_id_only_from_tool_messages():
    msgs = [Msg("human", "mi lead es L-AAAAAAAA"),
            Msg("tool", '{"status": "supported", "lead_id": "L-1234ABCD"}'),
            Msg("ai", "ok")]
    assert la.extraer_lead_id(msgs) == "L-1234ABCD"
    assert la.extraer_lead_id([Msg("human", "L-AAAAAAAA")]) is None


def test_event_detection():
    assert la.hubo_evento_accionable([Msg("tool", '{"status": "quotation_created", "quote_id": "Q-1"}')])
    assert la.hubo_evento_accionable([Msg("tool", '{"status": "scheduled", "type": "showroom_visit"}')])
    assert not la.hubo_evento_accionable([Msg("tool", '{"status": "supported", "lead_id": "L-1234ABCD"}')])
    assert not la.hubo_evento_accionable([Msg("ai", 'status "scheduled"')])  # only tool results count


def test_transcript_skips_tool_messages():
    msgs = [Msg("human", "Hola"), Msg("tool", "secret tool output"), Msg("ai", "Buenas")]
    t = la.construir_transcripcion(msgs)
    assert "Customer: Hola" in t and "Assistant: Buenas" in t and "secret" not in t


# --- end to end with a fake LLM ------------------------------------------------
def leer_scores(carpeta):
    ws = load_workbook(carpeta / "lead_scores.xlsx")["LeadScores"]
    cab = [c.value for c in ws[1]]
    return [dict(zip(cab, [c.value for c in f])) for f in ws.iter_rows(min_row=2)]


def test_analysis_saves_row_and_upserts(datos):
    llm = FakeLLM(SENALES_FUERTES)
    res = run(la.analizar_lead(LEAD, "Customer: necesito 4 pantallas para el 30", llm))
    assert res["score"] == 95 and res["priority"] == "HOT" and res["should_notify"] is True

    run(la.analizar_lead(LEAD, "Customer: otra vez", llm))   # re-analysis updates the same row
    filas = leer_scores(datos)
    assert len(filas) == 1 and filas[0]["priority"] == "HOT" and filas[0]["notified"] == "no"

    la.marcar_notificado(LEAD)
    res3 = run(la.analizar_lead(LEAD, "Customer: tercera", llm))
    assert res3["should_notify"] is False                    # no duplicate alerts
    assert leer_scores(datos)[0]["notified"] == "yes"


def test_llm_failure_falls_back_to_facts(datos):
    res = run(la.analizar_lead(LEAD, "x", FakeLLM(falla=True)))
    assert res["llm_ok"] is False
    assert res["score"] == 30 and res["priority"] == "COLD"  # 15 visit + 10 deal + 5 company
    assert leer_scores(datos)[0]["llm_ok"] == "no"


def test_contact_details_not_sent_to_llm(datos):
    llm = FakeLLM(SENALES_FUERTES)
    run(la.analizar_lead(LEAD, "Customer: hola", llm))
    enviado = str(llm.recibido)
    assert "ana@acme.com" not in enviado and "5512345678" not in enviado
    assert "untrusted" in enviado.lower()


def test_rules_file_is_created_and_editable(datos):
    la.asegurar_reglas()
    ruta = datos / "scoring_rules.xlsx"
    wb = load_workbook(ruta)
    for fila in wb["Thresholds"].iter_rows(min_row=2):
        if fila[0].value == "HOT":
            fila[1].value = 99          # the sales team makes HOT stricter
    wb.save(ruta)
    res = run(la.analizar_lead(LEAD, "x", FakeLLM(SENALES_FUERTES)))
    assert res["score"] == 95 and res["priority"] == "WARM"
